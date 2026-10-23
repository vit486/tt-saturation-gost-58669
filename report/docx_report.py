"""Генерация отчёта в формате Word (.docx), включая сводку по методам и графики."""
from __future__ import annotations

import io
import math
import os
import tempfile
from datetime import datetime

from docx import Document
from docx.shared import Inches

from methods.gost_58669 import (
    ANALYTIC_KEYS, GRAPHIC_KEYS, VAH_KEYS, BN_KEYS, K_R_PRED, ST_GRAPHIC_REQUIRED,
    ST_NO_SAT, ST_OK, ST_STEADY, METHOD_PRIORITY, build_kpr_curve, build_vah_default,
    build_bn_default, calc_b_t, method_enabled, summarize_method, determine_defining_method,
)

try:
    os.environ.setdefault("MPLCONFIGDIR", tempfile.mkdtemp(prefix="mpl_"))
    import matplotlib
    matplotlib.use("Agg")
    matplotlib.rcParams["font.family"] = "DejaVu Sans"
    _HAS_MPL = True
except Exception:
    _HAS_MPL = False

METHOD_TITLES = {
    "analytic": "аналитический метод (раздел 5.1)",
    "graphic": "графический метод по паспортным данным (раздел 5.2)",
    "both": "аналитический и графический методы (разделы 5.1 и 5.2)",
}


def _scenario_title(key, k_r):
    prefix = {
        "without_residual": "аналитический", "with_residual": "аналитический",
        "graphic_without_residual": "графический", "graphic_with_residual": "графический",
        "vah_without_residual": "по ВАХ", "vah_with_residual": "по ВАХ",
        "bn_without_residual": "по B(H)", "bn_with_residual": "по B(H)",
    }.get(key, key)
    suffix = ("без остаточной намагниченности (K_r = 0)" if "without" in key
              else f"с остаточной намагниченностью (K_r = {k_r:g})")
    return f"{prefix} — {suffix}"


def fmt_t(res) -> str:
    if res is None:
        return "—"
    if res.status == ST_OK and res.t_nas is not None:
        return f"{res.t_nas * 1000:.2f} мс"
    if res.status == ST_NO_SAT:
        return "насыщение отсутствует"
    if res.status == ST_STEADY:
        return "насыщение в установившемся режиме (A ≤ 1)"
    if res.status == ST_GRAPHIC_REQUIRED:
        return "метод неприменим — см. примечания"
    return "—"


def _add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = str(h)
    for row in rows:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    return table


def _keys_for(method: str):
    if method == "graphic":
        return GRAPHIC_KEYS
    if method == "both":
        return ANALYTIC_KEYS + GRAPHIC_KEYS
    return ANALYTIC_KEYS


def _method_rows(res, keys, k_r, name_prefix):
    rows = []
    for key in keys:
        r = res.get(key)
        if r is None:
            continue
        extra = ""
        if r.status == ST_OK and r.theta_star is not None:
            tag = "вручную" if r.theta_manual else "θ*"
            extra = f", {tag} = {math.degrees(r.theta_star):.1f}°"
        rows.append((f"{name_prefix} — {_scenario_title(key, k_r).split(' — ')[-1]}",
                     f"{r.a:.4f}", f"{fmt_t(r)}{extra}"))
    return rows


def _status_text(status, t_worst):
    if status is None:
        return "нет данных", "—"
    if status == "неприменим":
        return "неприменим (см. условия)", "—"
    if status == "недопустима (A ≤ 1)":
        return "A ≤ 1 — недопустима", "насыщение в уст. режиме"
    if status == "насыщения нет":
        return "насыщения нет", "∞"
    return "применим", (f"{t_worst * 1000:.2f} мс" if t_worst is not None else "—")


def _make_chart_png(kz, ct_m, gres, eff) -> io.BytesIO:
    if not _HAS_MPL:
        raise RuntimeError("matplotlib не установлен")
    import matplotlib.pyplot as plt
    if gres.status == ST_OK and gres.t_nas is not None:
        t_nas_ms = gres.t_nas * 1000.0
        t_end = min(max(t_nas_ms * 1.45, 20.0), 400.0)
    else:
        t_nas_ms = None
        t_end = min(max(kz.t_p * 3000.0, 40.0), 400.0)
    theta = gres.theta_star if gres.theta_star is not None else 0.0
    t_ms, k_star = build_kpr_curve(kz.t_p, ct_m.alpha, theta, t_end)
    _, k_zero = build_kpr_curve(kz.t_p, ct_m.alpha, 0.0, t_end)
    lbl_star = (f"θ = {math.degrees(theta):.1f}° (задано вручную)" if gres.theta_manual
                else "θ = θ* (расчётный минимум, п. Г.3.7)")
    fig, ax = plt.subplots(figsize=(8.0, 4.0))
    ax.plot(t_ms, k_star, color="#d62728", linewidth=2.0, label=lbl_star)
    ax.plot(t_ms, k_zero, color="#1f77b4", linewidth=1.3, label="θ = 0°")
    ax.axhline(eff, color="#2ca02c", linestyle="--", linewidth=1.5,
               label=f"Уровень A·(1−K_r) = {eff:.3f}")
    if t_nas_ms is not None:
        ax.axvline(t_nas_ms, color="#2ca02c", linestyle=":", linewidth=1.0)
        ax.plot([t_nas_ms], [eff], "o", color="#2ca02c", markersize=9,
                label=f"t_нас = {t_nas_ms:.2f} мс")
    ax.set_xlabel("Время t, мс")
    ax.set_ylabel("Коэффициент переходного режима K_п.р(t)")
    ax.set_title(f"{kz.label} · T_p = {kz.t_p * 1000:.0f} мс · cos α = {ct_m.cos_alpha:.3f}",
                 fontsize=10)
    ax.grid(True, alpha=0.3)
    ax.legend(fontsize=8, loc="best")
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150)
    plt.close(fig)
    buf.seek(0)
    return buf


def _make_vah_png(vah_data, r2) -> io.BytesIO:
    if not _HAS_MPL:
        raise RuntimeError("matplotlib не установлен")
    import matplotlib.pyplot as plt
    vah_plot = build_vah_default(vah_data['u_star'], vah_data['i0_star'], r2)
    fig, ax = plt.subplots(figsize=(8.0, 4.0))
    ax.plot(vah_plot['i0'], vah_plot['u'], color="#d62728", linewidth=2.0)
    ax.axhline(vah_data['u_star'], color="#2ca02c", linestyle="--", linewidth=1.2,
               label=f"Uε = {vah_data['u_star']:.1f} В")
    ax.axvline(vah_data['i0_star'], color="#1f77b4", linestyle=":", linewidth=1.2,
               label=f"I₀ε = {vah_data['i0_star']:.3f} А")
    ax.set_xlabel("Ток намагничивания I₀, А")
    ax.set_ylabel("Напряжение U, В")
    ax.set_title("Вольт-амперная характеристика ТТ (типовая)", fontsize=10)
    ax.grid(True, alpha=0.3)
    ax.legend(fontsize=8, loc="best")
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150)
    plt.close(fig)
    buf.seek(0)
    return buf


def _make_bn_png(bn_data) -> io.BytesIO:
    if not _HAS_MPL:
        raise RuntimeError("matplotlib не установлен")
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots(figsize=(8.0, 4.0))
    ax.plot(bn_data['h'], bn_data['b'], color="#9467bd", linewidth=2.0)
    ax.axhline(bn_data['b_sat'], color="#2ca02c", linestyle="--", linewidth=1.2,
               label=f"Bнас = {bn_data['b_sat']:.2f} Тл")
    ax.axhline(1.8, color="#ff7f0e", linestyle=":", linewidth=1.2, label="порог 1,8 Тл (усл. 21)")
    ax.set_xlabel("Напряжённость магнитного поля H, А/м")
    ax.set_ylabel("Магнитная индукция B, Тл")
    ax.set_title("Характеристика намагничивания B(H) (типовая)", fontsize=10)
    ax.grid(True, alpha=0.3)
    ax.legend(fontsize=8, loc="best")
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150)
    plt.close(fig)
    buf.seek(0)
    return buf


def build_report(ct, chain, results, verdict_lines, verdict, t_req_ms, method="analytic",
                 vah_data=None, bn_data=None, bn_params=None, k_r=K_R_PRED) -> bytes:
    """results — список кортежей (CTParams режима, KZParams, dict результатов)."""
    use_vah = vah_data is not None
    use_bn = bn_data is not None and bn_params is not None

    doc = Document()
    doc.add_heading("Отчёт о проверке трансформатора тока", 0)

    p = doc.add_paragraph()
    p.add_run("Дата формирования: ").bold = True
    p.add_run(datetime.now().strftime("%d.%m.%Y %H:%M"))
    doc.add_paragraph(
        "Определение времени до насыщения при коротких замыканиях по ГОСТ Р 58669-2019, "
        f"{METHOD_TITLES.get(method, method)}"
        + (" с привлечением метода по ВАХ (п. 5.3)" if use_vah else "")
        + (" и метода по характеристике намагничивания (п. 5.4)" if use_bn else "")
        + f". Расчёт выполнен без остаточной намагниченности и с K_r = {k_r:g} (п. 4.3)."
    )

    # --- Сводка по методам (от грубого к точному) ---
    def_name, def_keys, def_title = determine_defining_method(results, method, use_vah, use_bn)
    doc.add_heading("Сводка по методам (от грубого к точному)", level=1)
    sum_rows = []
    for mname, mkeys, mtitle in reversed(METHOD_PRIORITY):
        if not method_enabled(mname, method, use_vah, use_bn):
            sum_rows.append((mtitle, "не включён", "—", ""))
            continue
        status, t_worst = summarize_method(results, mkeys)
        st_text, t_text = _status_text(status, t_worst)
        sum_rows.append((mtitle, st_text, t_text, "✓ определяющий" if mname == def_name else ""))
    _add_table(doc, ["Метод", "Статус", "Наихудшее t_нас", "Роль"], sum_rows)
    doc.add_paragraph(f"Итоговая оценка вынесена по определяющему методу: {def_title}.")

    doc.add_heading("1. Исходные данные", level=1)
    _add_table(doc, ["Параметр", "Значение"], [
        ("Номинальный первичный ток I1ном", f"{ct.i1_nom:g} А"),
        ("Номинальный вторичный ток I2ном", f"{ct.i2_nom:g} А"),
        ("Номинальная предельная кратность Kном", f"{ct.k_nom:g}"),
        ("Номинальная мощность вторичной нагрузки Sном", f"{ct.s_nom:g} ВА"),
        ("Коэффициент мощности нагрузки cos φном", f"{ct.cos_phi_nom:g}"),
        ("Активное сопротивление вторичной обмотки R2", f"{ct.r2:g} Ом"),
        ("Индуктивное сопротивление вторичной обмотки X2", f"{ct.x2:g} Ом"),
        ("Сопротивление приборов Rприб", f"{chain['r_prib']:g} Ом"),
        ("Индуктивная составляющая нагрузки Xн", f"{chain['x_load']:g} Ом"),
        ("Схема соединения вторичных обмоток", chain["scheme"]),
        ("Материал жилы кабеля", chain["material"]),
        ("Длина кабеля (в одну сторону)", f"{chain['length']:g} м"),
        ("Сечение жилы", f"{chain['section']:g} мм²"),
        ("Переходное сопротивление контактов", f"{chain['r_per']:g} Ом"),
        ("Коэффициент остаточной намагниченности K_r", f"{k_r:g}"),
        ("Требуемое время до насыщения", f"{t_req_ms:g} мс" if t_req_ms else "не задано"),
    ])

    if vah_data:
        doc.add_paragraph(
            f"Метод по ВАХ (п. 5.3): Uε = {vah_data['u_star']:.1f} В, "
            f"I₀ε = {vah_data['i0_star']:.3f} А. Источник: паспорт ТТ / измерения по "
            "РД 153-34.0-35.301-2002, п. 3.7; для ТТ 35 кВ и выше ВАХ обязательна "
            "(ГОСТ Р 70507.2-2024, п. 6.7.5.2).")
    if use_bn:
        doc.add_paragraph(
            f"Метод по характеристике намагничивания (п. 5.4): w2 = {bn_params['w2']:g}, "
            f"s = {bn_params['s'] * 1e4:.1f} см², l = {bn_params['l']:g} м, "
            f"Bнас = {bn_data['b_sat']:.2f} Тл, Hколено = {bn_data['h_knee']:.0f} А/м. "
            "Источник: паспорт / заводская документация ТТ (пример — ГОСТ Р 58669-2019, В.4.1.1).")

    doc.add_heading("2. Промежуточные величины (общие)", level=1)
    _add_table(doc, ["Величина", "Значение"], [
        ("Сопротивление номинальной нагрузки zн.ном = Sном / I2ном²", f"{ct.z_n_nom:.4f} Ом"),
        ("Номинальное полное сопротивление z2Σном (формула 10)", f"{ct.z2_sigma_nom:.4f} Ом"),
        ("Коэффициент расчётной длины кабеля — трёхфазное КЗ", f"{chain['k_sch_3']:g}"),
        ("Коэффициент расчётной длины кабеля — однофазное КЗ", f"{chain['k_sch_1']:g}"),
    ])

    keys = _keys_for(method)
    for idx, (ct_m, kz, res) in enumerate(results, start=3):
        r_cable = ct_m.r_load_fact - chain["r_prib"] - chain["r_per"]
        doc.add_heading(f"{idx}. Режим: {kz.label}", level=1)
        _add_table(doc, ["Параметр", "Значение"], [
            ("Ток КЗ (действующее значение)", f"{kz.i_kz:g} А"),
            ("Постоянная времени T_p", f"{kz.t_p:g} с ({kz.t_p * 1000:.0f} мс)"),
            ("Фактическая кратность Kфакт = Iкз / I1ном", f"{kz.i_kz / ct_m.i1_nom:.2f}"),
            ("R кабеля для этого режима (ρ·k·L/S)", f"{r_cable:.4f} Ом"),
            ("Rн.факт = Rприб + Rпер + Rкаб", f"{ct_m.r_load_fact:.4f} Ом"),
            ("z2Σ (формула 11)", f"{ct_m.z2_sigma:.4f} Ом"),
            ("Угол α (формула 12)",
             f"{math.degrees(ct_m.alpha):.2f}° (cos α = {ct_m.cos_alpha:.4f})"),
            ("Параметр режима A (формула 9)", f"{res[keys[0]].a:.4f}"),
        ])

        if method == "both":
            rows = _method_rows(res, ANALYTIC_KEYS, k_r, "аналитический") + \
                   _method_rows(res, GRAPHIC_KEYS, k_r, "графический")
        else:
            rows = _method_rows(res, keys, k_r,
                                "аналитический" if method == "analytic" else "графический")
        _add_table(doc, ["Метод и сценарий", "Параметр A", "Время до насыщения"], rows)

        if use_vah and VAH_KEYS[0] in res:
            doc.add_paragraph("Результаты по ВАХ (п. 5.3):").bold = True
            _add_table(doc, ["Метод и сценарий", "Параметр A", "Время до насыщения"],
                       _method_rows(res, VAH_KEYS, k_r, "по ВАХ"))

        if use_bn and BN_KEYS[0] in res:
            b_t = calc_b_t(ct_m, kz, bn_params['w2'], bn_params['s'])
            doc.add_paragraph(
                f"Результаты по характеристике намагничивания (п. 5.4), Вт = {b_t:.3f} Тл:").bold = True
            _add_table(doc, ["Метод и сценарий", "Параметр A", "Время до насыщения"],
                       _method_rows(res, BN_KEYS, k_r, "по B(H)"))

        seen = set()
        all_keys = list(keys)
        if use_vah and VAH_KEYS[0] in res:
            all_keys += list(VAH_KEYS)
        if use_bn and BN_KEYS[0] in res:
            all_keys += list(BN_KEYS)
        for key in all_keys:
            r = res.get(key)
            if r is None:
                continue
            for note in r.notes:
                if note not in seen:
                    seen.add(note)
                    doc.add_paragraph(note, style="List Bullet")

        if method in ("graphic", "both"):
            for key in GRAPHIC_KEYS:
                r = res.get(key)
                if r is None or r.status != ST_OK or r.t_nas is None:
                    continue
                eff = r.a * (1.0 - r.k_r)
                cap = doc.add_paragraph()
                cap.add_run(f"График K_п.р(t) — {_scenario_title(key, k_r).split(' — ')[-1]} "
                            f"(θ* = {math.degrees(r.theta_star):.1f}°):").bold = True
                try:
                    doc.add_picture(_make_chart_png(kz, ct_m, r, eff), width=Inches(6.0))
                except Exception as exc:
                    doc.add_paragraph(f"(График не построен: {type(exc).__name__}.)")

    if vah_data:
        cap = doc.add_paragraph()
        cap.add_run("Вольт-амперная характеристика ТТ (типовая):").bold = True
        try:
            doc.add_picture(_make_vah_png(vah_data, ct.r2), width=Inches(6.0))
        except Exception as exc:
            doc.add_paragraph(f"(График ВАХ не построен: {type(exc).__name__}.)")

    if use_bn:
        cap = doc.add_paragraph()
        cap.add_run("Характеристика намагничивания B(H) (типовая):").bold = True
        try:
            doc.add_picture(_make_bn_png(bn_data), width=Inches(6.0))
        except Exception as exc:
            doc.add_paragraph(f"(График B(H) не построен: {type(exc).__name__}.)")

    doc.add_heading("Вывод", level=1)
    label = {"success": "ТТ проходит проверку.",
             "warning": "ТТ проходит проверку с замечаниями.",
             "danger": "ТТ проверку не проходит."}.get(verdict, "")
    if label:
        p = doc.add_paragraph()
        p.add_run(label).bold = True
    for line in verdict_lines:
        doc.add_paragraph(line, style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# === КОНЕЦ ФАЙЛА report/docx_report.py ===
