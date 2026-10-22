"""
⚡ Проверка трансформаторов тока на насыщение при коротких замыканиях.
Методика: ГОСТ Р 58669-2019, раздел 5.1 (аналитический метод).
Релиз 1.1: исправлена формула времени до насыщения (числитель ω·T_p),
добавлена самопроверка по эталонным примерам Приложения В ГОСТ Р 58669-2019.
"""

import io
import math
from dataclasses import dataclass, field
from datetime import datetime
from typing import List, Optional

import streamlit as st
from docx import Document

# ===========================================================================
# 1. РАСЧЁТНОЕ ЯДРО — ГОСТ Р 58669-2019, аналитический метод (п. 5.1)
# ===========================================================================

OMEGA = 2.0 * math.pi * 50.0   # угловая частота, рад/с (f = 50 Гц)
K_R_PRED = 0.86                # предельный коэффициент остаточной намагниченности (п. 5.1.3)
T_NAS_MIN_MS = 15.0            # порог применимости аналитического метода (п. 5.1.7)
T_P_RANGE_S = (0.03, 0.30)     # диапазон Tp универсальных характеристик (приложение Б)


@dataclass
class CTParams:
    """Паспортные данные трансформатора тока."""
    i1_nom: float          # номинальный первичный ток, А
    i2_nom: float          # номинальный вторичный ток, А
    k_nom: float           # номинальная предельная кратность
    s_nom: float           # номинальная мощность вторичной нагрузки, ВА
    cos_phi_nom: float     # коэффициент мощности номинальной нагрузки
    r2: float              # активное сопротивление вторичной обмотки, Ом
    x2: float              # индуктивное сопротивление вторичной обмотки, Ом
    r_load_fact: float     # фактическое активное сопротивление нагрузки, Ом
    x_load_fact: float     # фактическое индуктивное сопротивление нагрузки, Ом

    @property
    def z_n_nom(self) -> float:
        """Полное сопротивление номинальной нагрузки, Ом."""
        return self.s_nom / (self.i2_nom ** 2)

    @property
    def z2_sigma_nom(self) -> float:
        """Номинальное полное сопротивление вторичной нагрузки ТТ, Ом (формула 10)."""
        sin_phi = math.sqrt(max(0.0, 1.0 - self.cos_phi_nom ** 2))
        return math.sqrt(
            (self.r2 + self.z_n_nom * self.cos_phi_nom) ** 2
            + (self.x2 + self.z_n_nom * sin_phi) ** 2
        )

    @property
    def z2_sigma(self) -> float:
        """Полное сопротивление ветви вторичного тока, Ом (формула 11)."""
        return math.sqrt(
            (self.r2 + self.r_load_fact) ** 2
            + (self.x2 + self.x_load_fact) ** 2
        )


@dataclass
class KZParams:
    """Параметры короткого замыкания."""
    i_kz: float            # действующее значение тока КЗ (первичная сторона), А
    t_p: float             # постоянная времени затухания апериодической составляющей, с
    label: str = "трёхфазное КЗ"


@dataclass
class SaturationResult:
    """Результат для одного режима КЗ и одного значения остаточной намагниченности."""
    a: float                            # параметр режима, формула (9)
    k_r: float                          # коэффициент остаточной намагниченности
    status: str                         # ok | no_saturation | steady_saturation | graphic_required
    t_nas: Optional[float]              # время до насыщения, с (только при status == "ok")
    notes: List[str] = field(default_factory=list)


def _calc_one_mode(ct: CTParams, kz: KZParams, k_r: float) -> SaturationResult:
    """Расчёт времени до насыщения для заданного коэффициента остаточной намагниченности."""
    notes: List[str] = []

    if kz.i_kz <= 0:
        raise ValueError("Ток КЗ должен быть больше нуля.")
    if ct.z2_sigma == 0:
        raise ValueError(
            "Полное сопротивление вторичной цепи равно нулю — проверьте исходные данные."
        )

    # Параметр режима A — формула (9)
    a = (ct.i1_nom * ct.k_nom * ct.z2_sigma_nom) / (kz.i_kz * ct.z2_sigma)

    # Приложение Б: универсальные характеристики построены для Tp = 30–300 мс
    if not (T_P_RANGE_S[0] <= kz.t_p <= T_P_RANGE_S[1]):
        notes.append(
            f"T_p = {kz.t_p * 1000:.0f} мс выходит за диапазон "
            f"{T_P_RANGE_S[0] * 1000:.0f}–{T_P_RANGE_S[1] * 1000:.0f} мс, "
            "для которого построены универсальные характеристики (приложение Б)."
        )

    effective = a * (1.0 - k_r)

    # Условие (5): A > 1 — только для варианта без остаточной намагниченности
    if k_r == 0.0 and a <= 1.0:
        notes.append(
            "Условие (5) ГОСТ не выполняется: A ≤ 1 — действующий ток КЗ превышает "
            "ток предельной кратности ТТ. Эксплуатация ТТ в таких условиях недопустима."
        )
        return SaturationResult(a, k_r, "steady_saturation", 0.0, notes)

    # Условие (8): A*(1-Kr) > 1, иначе формула даёт отрицательное время
    if effective <= 1.0:
        notes.append(
            "Условие (8) ГОСТ не выполняется: A·(1−K_r) ≤ 1 — аналитический метод "
            "неприменим. Согласно п. 5.1.5 следует использовать графический метод "
            "по паспортным данным или по ВАХ."
        )
        return SaturationResult(a, k_r, "graphic_required", None, notes)

    # Условие (7)/(4): если не выполняется — насыщения нет
    numerator = OMEGA * kz.t_p
    denom = numerator + 1.0 - effective
    if denom <= 0.0:
        notes.append(
            "Условие (7) ГОСТ не выполняется: ω·T_p + 1 ≤ A·(1−K_r) — "
            "насыщение магнитопровода в переходном процессе отсутствует."
        )
        return SaturationResult(a, k_r, "no_saturation", None, notes)

    # Формулы (3)/(6) ГОСТ Р 58669-2019:
    # t_нас = T_p * ln( ω*T_p / (ω*T_p + 1 - A*(1-K_r)) )
    t_nas = kz.t_p * math.log(numerator / denom)

    # Пункт 5.1.7
    if t_nas * 1000.0 < T_NAS_MIN_MS:
        notes.append(
            f"t_нас = {t_nas * 1000:.1f} мс < {T_NAS_MIN_MS:.0f} мс — согласно п. 5.1.7 "
            "результат следует уточнить графическим методом по паспортным данным."
        )

    return SaturationResult(a, k_r, "ok", t_nas, notes)


def calculate(ct: CTParams, kz: KZParams) -> dict:
    """Расчёт для обоих вариантов остаточной намагниченности (п. 4.3 ГОСТ)."""
    return {
        "without_residual": _calc_one_mode(ct, kz, 0.0),
        "with_residual": _calc_one_mode(ct, kz, K_R_PRED),
    }


# ===========================================================================
# 2. ГЕНЕРАЦИЯ ОТЧЁТА WORD
# ===========================================================================

def fmt_t(res: SaturationResult) -> str:
    if res.status == "ok":
        return f"{res.t_nas * 1000:.1f} мс"
    if res.status == "no_saturation":
        return "насыщение отсутствует"
    if res.status == "steady_saturation":
        return "насыщение в установившемся режиме"
    return "требуется графический метод"


def _add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = str(h)
    for row in rows:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    return table


def build_report(ct, results, verdict_lines, verdict, t_req_ms) -> bytes:
    """Формирует байты .docx-файла."""
    doc = Document()
    doc.add_heading("Отчёт о проверке трансформатора тока", 0)

    p = doc.add_paragraph()
    p.add_run("Дата формирования: ").bold = True
    p.add_run(datetime.now().strftime("%d.%m.%Y %H:%M"))
    doc.add_paragraph(
        "Определение времени до насыщения аналитическим методом по ГОСТ Р 58669-2019 "
        "(раздел 5.1). Расчёт выполнен без остаточной намагниченности и с предельной "
        f"остаточной намагниченностью K_r = {K_R_PRED} (п. 4.3, п. 5.1.3)."
    )

    doc.add_heading("1. Исходные данные", level=1)
    _add_table(doc, ["Параметр", "Значение"], [
        ("Номинальный первичный ток I1ном", f"{ct.i1_nom:g} А"),
        ("Номинальный вторичный ток I2ном", f"{ct.i2_nom:g} А"),
        ("Номинальная предельная кратность Kном", f"{ct.k_nom:g}"),
        ("Номинальная мощность вторичной нагрузки Sном", f"{ct.s_nom:g} ВА"),
        ("Коэффициент мощности нагрузки cos φном", f"{ct.cos_phi_nom:g}"),
        ("Активное сопротивление вторичной обмотки R2", f"{ct.r2:g} Ом"),
        ("Индуктивное сопротивление вторичной обмотки X2", f"{ct.x2:g} Ом"),
        ("Фактическое активное сопротивление нагрузки Rн.факт", f"{ct.r_load_fact:g} Ом"),
        ("Фактическое индуктивное сопротивление нагрузки Xн.факт", f"{ct.x_load_fact:g} Ом"),
        ("Требуемое время до насыщения", f"{t_req_ms:g} мс" if t_req_ms else "не задано"),
    ])

    doc.add_heading("2. Промежуточные величины", level=1)
    _add_table(doc, ["Величина", "Значение"], [
        ("Сопротивление номинальной нагрузки zн.ном = Sном / I2ном²", f"{ct.z_n_nom:.4f} Ом"),
        ("Номинальное полное сопротивление z2Σном (формула 10)", f"{ct.z2_sigma_nom:.4f} Ом"),
        ("Полное сопротивление ветви вторичного тока z2Σ (формула 11)", f"{ct.z2_sigma:.4f} Ом"),
    ])

    for idx, (kz, res) in enumerate(results, start=3):
        doc.add_heading(f"{idx}. Режим: {kz.label}", level=1)
        _add_table(doc, ["Параметр", "Значение"], [
            ("Ток КЗ (действующее значение)", f"{kz.i_kz:g} А"),
            ("Постоянная времени T_p", f"{kz.t_p:g} с ({kz.t_p * 1000:.0f} мс)"),
            ("Фактическая кратность Kфакт = Iкз / I1ном", f"{kz.i_kz / ct.i1_nom:.2f}"),
        ])
        _add_table(doc, ["Сценарий", "Параметр A", "Время до насыщения"], [
            ("Без остаточной намагниченности (K_r = 0)",
             f"{res['without_residual'].a:.3f}", fmt_t(res["without_residual"])),
            (f"С предельной остаточной намагниченностью (K_r = {K_R_PRED})",
             f"{res['with_residual'].a:.3f}", fmt_t(res["with_residual"])),
        ])
        for key in ("without_residual", "with_residual"):
            for note in res[key].notes:
                doc.add_paragraph(note, style="List Bullet")

    doc.add_heading("Вывод", level=1)
    label = {
        "success": "ТТ проходит проверку.",
        "warning": "ТТ проходит проверку с замечаниями.",
        "danger": "ТТ проверку не проходит.",
    }.get(verdict, "")
    if label:
        p = doc.add_paragraph()
        p.add_run(label).bold = True
    for line in verdict_lines:
        doc.add_paragraph(line, style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ===========================================================================
# 3. ЭТАЛОННЫЕ ПРИМЕРЫ ДЛЯ САМОПРОВЕРКИ
#    Источник: ГОСТ Р 58669-2019, Приложение В (В.1.1–В.1.3), с учётом
#    Поправки (ИУС 6–7, 2020).
#    Примечание: в примере В.1.2 в тексте стандарта опечатка — в исходных
#    данных указан Kном=20, однако расчёт (В.1.2.3, В.2.2.1) выполнен при
#    Kном=18; в эталоне принято значение 18.
# ===========================================================================

REFERENCE_EXAMPLES = [
    {
        "id": "В.1.1",
        "name": "SAS 550/5G 2000/1",
        "ct": dict(i1_nom=2000.0, i2_nom=1.0, k_nom=20.0, s_nom=40.0,
                   cos_phi_nom=0.8, r2=7.51, x2=0.0),
        "modes": [
            {"label": "трёхфазное КЗ", "i_kz": 23145.0, "t_p": 0.128,
             "r_load": 6.3, "x_load": 0.0,
             "a_expected": 5.785, "t_expected_ms": 16.22,
             "residual_expected": "graphic_required"},
            {"label": "однофазное КЗ", "i_kz": 26900.0, "t_p": 0.180,
             "r_load": 12.6, "x_load": 0.0,
             "a_expected": 3.418, "t_expected_ms": 7.87,
             "residual_expected": "graphic_required"},
        ],
    },
    {
        "id": "В.1.2",
        "name": "ТФЗМ-500Б-1У 2000/1",
        "ct": dict(i1_nom=2000.0, i2_nom=1.0, k_nom=18.0, s_nom=50.0,
                   cos_phi_nom=0.8, r2=10.0, x2=0.0),
        "modes": [
            {"label": "трёхфазное КЗ", "i_kz": 23145.0, "t_p": 0.128,
             "r_load": 6.3, "x_load": 0.0,
             "a_expected": 5.564, "t_expected_ms": 15.42,
             "residual_expected": "graphic_required"},
            {"label": "однофазное КЗ", "i_kz": 26900.0, "t_p": 0.180,
             "r_load": 12.6, "x_load": 0.0,
             "a_expected": 3.453, "t_expected_ms": 7.98,
             "residual_expected": "graphic_required"},
        ],
    },
    {
        "id": "В.1.3",
        "name": "JK ELK СВЗ 4000/1",
        "ct": dict(i1_nom=4000.0, i2_nom=1.0, k_nom=20.0, s_nom=30.0,
                   cos_phi_nom=0.8, r2=10.5, x2=0.0),
        "modes": [
            {"label": "трёхфазное КЗ", "i_kz": 23145.0, "t_p": 0.128,
             "r_load": 6.3, "x_load": 0.0,
             "a_expected": 8.0, "t_expected_ms": 24.5,
             "residual_expected": "ok", "t_residual_expected_ms": 0.39},
            {"label": "однофазное КЗ", "i_kz": 26900.0, "t_p": 0.180,
             "r_load": 12.6, "x_load": 0.0,
             "a_expected": 5.0, "t_expected_ms": 13.24,
             "residual_expected": "graphic_required"},
        ],
    },
]

A_TOL = 0.02      # допуск на A (в стандарте значения округлены до 3 знаков)
T_TOL_MS = 0.2    # допуск на время до насыщения, мс


def run_self_checks():
    """Прогон расчётов по эталонным примерам ГОСТ Р 58669-2019 (Приложение В)."""
    rows = []
    all_ok = True
    for ex in REFERENCE_EXAMPLES:
        for mode in ex["modes"]:
            ct = CTParams(
                **ex["ct"],
                r_load_fact=mode["r_load"],
                x_load_fact=mode["x_load"],
            )
            kz = KZParams(mode["i_kz"], mode["t_p"], mode["label"])
            res = calculate(ct, kz)

            r0 = res["without_residual"]
            a_ok = abs(r0.a - mode["a_expected"]) <= A_TOL
            t_ok = (r0.status == "ok"
                    and abs(r0.t_nas * 1000.0 - mode["t_expected_ms"]) <= T_TOL_MS)

            r_res = res["with_residual"]
            if mode["residual_expected"] == "graphic_required":
                res_ok = r_res.status == "graphic_required"
                res_text = "графический метод (условие (8) не выполняется)"
            else:
                res_ok = (r_res.status == "ok" and
                          abs(r_res.t_nas * 1000.0 - mode["t_residual_expected_ms"]) <= T_TOL_MS)
                res_text = (f"{r_res.t_nas * 1000.0:.2f} мс "
                            f"(эталон {mode['t_residual_expected_ms']} мс)"
                            if r_res.status == "ok" else r_res.status)

            ok = a_ok and t_ok and res_ok
            all_ok = all_ok and ok
            rows.append({
                "Пример": f"{ex['id']} {ex['name']}",
                "Режим": mode["label"],
                "A эталон": mode["a_expected"],
                "A расчёт": round(r0.a, 3),
                "t эталон, мс": mode["t_expected_ms"],
                "t расчёт, мс": round(r0.t_nas * 1000.0, 2) if r0.status == "ok" else "—",
                "Сценарий K_r=0,86": res_text,
                "Статус": "✅" if ok else "❌",
            })
    return rows, all_ok


# ===========================================================================
# 4. ИНТЕРФЕЙС
# ===========================================================================

st.set_page_config(
    page_title="Проверка ТТ — ГОСТ Р 58669-2019",
    layout="wide",
    page_icon="⚡",
)

st.title("⚡ Проверка трансформаторов тока")
st.caption("Насыщение при КЗ · ГОСТ Р 58669-2019 · аналитический метод · релиз 1.1")

# Пресеты ТТ — задел под будущую базу типовых трансформаторов
PRESETS = {
    "— ввести вручную —": {},
    "Пример: ТОЛ-10 600/5 10P10 15 ВА": dict(
        i1_nom=600.0, i2_nom=5.0, k_nom=10.0, s_nom=15.0, cos_phi_nom=0.8,
        r2=0.12, x2=0.0, r_load_fact=0.35, x_load_fact=0.0,
    ),
    "Пример: ТФЗМ-110 2000/5 10P20 30 ВА": dict(
        i1_nom=2000.0, i2_nom=5.0, k_nom=20.0, s_nom=30.0, cos_phi_nom=0.8,
        r2=0.45, x2=0.0, r_load_fact=0.8, x_load_fact=0.0,
    ),
}

with st.sidebar:
    st.header("⚙️ Исходные данные")

    preset_name = st.selectbox("Тип ТТ", list(PRESETS.keys()))
    preset = PRESETS[preset_name]

    st.subheader("Паспортные данные ТТ")
    i1_nom = st.number_input(
        "I1ном, А", min_value=1.0, value=float(preset.get("i1_nom", 1000.0)),
        step=50.0, key=f"{preset_name}_i1")
    i2_nom = st.selectbox(
        "I2ном, А", (1.0, 2.0, 5.0),
        index=(1.0, 2.0, 5.0).index(float(preset.get("i2_nom", 5.0))),
        key=f"{preset_name}_i2")
    k_nom = st.number_input(
        "Kном (предельная кратность)", min_value=1.0,
        value=float(preset.get("k_nom", 10.0)), step=1.0, key=f"{preset_name}_k")
    s_nom = st.number_input(
        "Sном вторичной нагрузки, ВА", min_value=1.0,
        value=float(preset.get("s_nom", 15.0)), step=1.0, key=f"{preset_name}_s")
    cos_phi = st.slider(
        "cos φ ном. нагрузки", 0.1, 1.0,
        float(preset.get("cos_phi_nom", 0.8)), 0.05, key=f"{preset_name}_cos")
    r2 = st.number_input(
        "R2 вторичной обмотки, Ом", min_value=0.0,
        value=float(preset.get("r2", 0.15)), step=0.01, format="%.3f",
        key=f"{preset_name}_r2")
    x2 = st.number_input(
        "X2 вторичной обмотки, Ом", min_value=0.0,
        value=float(preset.get("x2", 0.0)), step=0.01, format="%.3f",
        key=f"{preset_name}_x2")

    st.subheader("Фактическая нагрузка")
    r_load = st.number_input(
        "Rн.факт, Ом", min_value=0.0,
        value=float(preset.get("r_load_fact", 0.5)), step=0.01, format="%.3f",
        key=f"{preset_name}_rl")
    x_load = st.number_input(
        "Xн.факт, Ом", min_value=0.0,
        value=float(preset.get("x_load_fact", 0.0)), step=0.01, format="%.3f",
        key=f"{preset_name}_xl")

    st.subheader("Короткое замыкание")
    i_kz3 = st.number_input("Iкз(3), А", min_value=1.0, value=15000.0, step=100.0)
    tp3 = st.number_input(
        "Tр.экв для трёхфазного КЗ, с", min_value=0.001, value=0.05,
        step=0.005, format="%.3f")
    st.caption("Например: 0.05 с = 50 мс")

    use_kz1 = st.checkbox("Добавить расчёт для однофазного КЗ (п. 4.2.4 ГОСТ)")
    i_kz1 = tp1 = None
    if use_kz1:
        i_kz1 = st.number_input("Iкз(1), А", min_value=1.0, value=12000.0, step=100.0)
        tp1 = st.number_input(
            "Tр.экв для однофазного КЗ, с", min_value=0.001, value=0.05,
            step=0.005, format="%.3f")

    st.subheader("Критерий проверки")
    t_req = st.number_input(
        "Требуемое t_нас, мс (0 — не проверять)", 0.0, 500.0, 0.0, 5.0,
        help="Минимально необходимое время до насыщения по требованиям РЗА "
             "(например, время отключения КЗ с запасом).")

# ---------------------------------------------------------------------------
# Расчёт
# ---------------------------------------------------------------------------
ct = CTParams(i1_nom, i2_nom, k_nom, s_nom, cos_phi, r2, x2, r_load, x_load)

kz_list = [KZParams(i_kz3, tp3, "трёхфазное КЗ")]
if use_kz1:
    kz_list.append(KZParams(i_kz1, tp1, "однофазное КЗ"))

try:
    results = [(kz, calculate(ct, kz)) for kz in kz_list]
except ValueError as e:
    st.error(f"Ошибка исходных данных: {e}")
    st.stop()


def build_verdict(res_list, t_req_ms):
    """Итоговый вывод: минимум по всем режимам и сценариям (п. 4.2.4 ГОСТ)."""
    lines = []
    steady = any(r.status == "steady_saturation" for _, res in res_list for r in res.values())
    t_values = [r.t_nas for _, res in res_list for r in res.values() if r.status == "ok"]
    graphic_needed = any(r.status == "graphic_required"
                         for _, res in res_list for r in res.values())

    if steady:
        verdict = "danger"
        lines.append("A ≤ 1: ток КЗ превышает ток предельной кратности ТТ — насыщение "
                     "наступает в установившемся режиме. Эксплуатация недопустима.")
        lines.append("Рекомендация: выбрать ТТ с большей Kном или большим коэффициентом "
                     "трансформации, либо уменьшить вторичную нагрузку.")
    elif not t_values:
        if graphic_needed:
            verdict = "warning"
            lines.append("Аналитический метод неприменим ни в одном сценарии — "
                         "требуется графический метод (п. 5.1.5 ГОСТ).")
        else:
            verdict = "success"
            lines.append("Во всех расчётных режимах насыщение магнитопровода отсутствует.")
    else:
        t_min_ms = min(t_values) * 1000.0
        lines.append(f"Наименьшее время до насыщения по всем режимам и сценариям: "
                     f"{t_min_ms:.1f} мс.")
        if t_req_ms > 0:
            if t_min_ms >= t_req_ms:
                verdict = "success"
                lines.append(f"Требование выполняется: {t_min_ms:.1f} мс ≥ {t_req_ms:g} мс.")
            else:
                verdict = "danger"
                lines.append(f"Требование НЕ выполняется: {t_min_ms:.1f} мс < {t_req_ms:g} мс.")
                lines.append("Рекомендация: выбрать ТТ с большей Kном или Sном, либо "
                             "уменьшить сопротивление вторичной цепи.")
        else:
            verdict = "success"
        if verdict == "success" and graphic_needed:
            verdict = "warning"
            lines.append("Для сценария с остаточной намагниченностью аналитический метод "
                         "неприменим (п. 5.1.5) — при необходимости уточните результат "
                         "графическим методом.")
    return verdict, lines


verdict, verdict_lines = build_verdict(results, t_req)

# ---------------------------------------------------------------------------
# Вкладки
# ---------------------------------------------------------------------------
tab_calc, tab_check, tab_theory, tab_about = st.tabs([
    "🧮 Результаты расчёта",
    "🧪 Самопроверка (эталонные примеры)",
    "📚 Теория и формулы (ГОСТ Р 58669-2019)",
    "ℹ️ О приложении и планы",
])

with tab_calc:
    banner = {"success": st.success, "warning": st.warning, "danger": st.error}[verdict]
    banner(" ".join(verdict_lines))

    for kz, res in results:
        st.subheader(f"Режим: {kz.label}")
        c1, c2, c3 = st.columns(3)
        c1.metric("Iкз", f"{kz.i_kz:g} А")
        c2.metric("Kфакт = Iкз / I1ном", f"{kz.i_kz / ct.i1_nom:.2f}")
        c3.metric("Tр.экв", f"{kz.t_p * 1000:.0f} мс")

        rows = []
        for name, key in (("Без остаточной намагниченности (K_r = 0)", "without_residual"),
                          (f"С предельной остаточной намагниченностью (K_r = {K_R_PRED})",
                           "with_residual")):
            r = res[key]
            if r.status == "ok":
                t_str = f"{r.t_nas * 1000:.1f} мс"
            elif r.status == "no_saturation":
                t_str = "∞ — насыщения нет"
            elif r.status == "steady_saturation":
                t_str = "0 мс — насыщение в установившемся режиме"
            else:
                t_str = "требуется графический метод"
            rows.append({"Сценарий": name, "Параметр A": round(r.a, 3),
                         "Время до насыщения": t_str})
        st.table(rows)

        for key in ("without_residual", "with_residual"):
            for note in res[key].notes:
                st.caption(f"• {note}")

    with st.expander("🔍 Промежуточные величины"):
        st.markdown(
            f"- zн.ном = Sном / I2ном² = **{ct.z_n_nom:.4f} Ом**\n"
            f"- z2Σном (формула 10) = **{ct.z2_sigma_nom:.4f} Ом**\n"
            f"- z2Σ (формула 11) = **{ct.z2_sigma:.4f} Ом**\n"
            f"- ω = 2π·50 = **314.16 рад/с**"
        )

    st.divider()
    report_bytes = build_report(ct, results, verdict_lines, verdict, t_req)
    st.download_button(
        "📄 Скачать отчёт (.docx)",
        data=report_bytes,
        file_name=f"Otchet_TT_GOST_58669_{datetime.now():%Y%m%d_%H%M}.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

with tab_check:
    st.markdown(
        "Автоматическая сверка расчёта с **эталонными примерами ГОСТ Р 58669-2019 "
        "(Приложение В, В.1.1–В.1.3)** с учётом Поправки (ИУС 6–7, 2020). "
        "Допуски: параметр A ±0,02; время ±0,2 мс (в стандарте значения округлены). "
        "Если хотя бы одна строка ❌ — результатам расчёта доверять нельзя."
    )
    check_rows, all_ok = run_self_checks()
    if all_ok:
        st.success("✅ Все эталонные примеры воспроизведены — расчётные формулы "
                   "соответствуют ГОСТ Р 58669-2019.")
    else:
        st.error("❌ Обнаружено расхождение с эталонными примерами. Пользоваться "
                 "результатами расчёта до устранения не следует.")
    st.dataframe(check_rows, hide_index=True)
    st.caption(
        "Примечание: в примере В.1.2 в тексте стандарта опечатка (в исходных данных "
        "Kном=20, расчёт выполнен при Kном=18) — в эталоне принято Kном=18."
    )

with tab_theory:
    st.markdown("### Аналитический метод (раздел 5.1 ГОСТ Р 58669-2019)")
    st.markdown("**1. Сопротивления вторичной цепи.** Номинальная нагрузка: "
                "`zн.ном = Sном / I2ном²`. Номинальное полное сопротивление "
                "вторичной нагрузки ТТ (формула 10):")
    st.latex(r"z_{2\Sigma\,nom} = \sqrt{(R_2 + z_{n\,nom}\cos\varphi_{nom})^2 + "
             r"(X_2 + z_{n\,nom}\sin\varphi_{nom})^2}")
    st.markdown("Полное сопротивление ветви вторичного тока (формула 11):")
    st.latex(r"z_{2\Sigma} = \sqrt{(R_2 + R_{fact})^2 + (X_2 + X_{fact})^2}")

    st.markdown("**2. Параметр режима A (формула 9)** — отношение потокосцепления "
                "насыщения к амплитуде потокосцепления от периодической составляющей тока КЗ:")
    st.latex(r"A = \frac{I_{1nom} \cdot K_{nom} \cdot z_{2\Sigma\,nom}}"
             r"{I_{KZ} \cdot z_{2\Sigma}}")

    st.markdown("**3. Время до насыщения** — без остаточной намагниченности "
                "(формула 3) и с предельной остаточной намагниченностью "
                f"K_r = {K_R_PRED} (формула 6):")
    st.latex(r"t_{nas} = T_p \cdot \ln\frac{\omega T_p}"
             r"{\omega T_p + 1 - A(1 - K_r)}")

    st.markdown("""
**4. Проверяемые условия**

| Условие | Смысл | Действие при нарушении |
|---|---|---|
| A > 1 (условие 5) | ток предельной кратности больше тока КЗ | эксплуатация ТТ недопустима |
| A·(1−K_r) > 1 (условие 8) | применимость формулы времени | графический метод (п. 5.1.5) |
| ω·T_p + 1 > A·(1−K_r) (условие 7) | наличие насыщения | насыщения нет, t_нас = ∞ |
| t_нас ≥ 15 мс (п. 5.1.7) | точность аналитического метода | уточнить графическим методом |

Расчёт выполняется для трёхфазного и (в сетях с глухозаземлённой нейтралью) однофазного КЗ;
итогом является **наименьшее** из полученных значений (п. 4.2.4).
""")

with tab_about:
    st.markdown("""
Приложение выполняет проверку трансформаторов тока на насыщение при коротких замыканиях
в соответствии с **ГОСТ Р 58669-2019** (аналитический метод, раздел 5.1).

**Журнал изменений:**
- **Релиз 1.1 (22.10.2026):** исправлена формула времени до насыщения — числитель
  ω·T_p, а не ω·T_p+1 (ошибка выявлена сверкой с эталонными примерами Приложения В;
  старая формула завышала время до насыщения, т. е. давала неконсервативный результат).
  Добавлена вкладка «Самопроверка».

**Планы развития:**
- релиз 1.2 — модульная структура (каждая методика — отдельный файл);
- релиз 1.3 — графические методы ГОСТ Р 58669-2019 (по паспортным данным, ВАХ,
  характеристике намагничивания) с динамическими графиками;
- релиз 1.4 — ГОСТ Р 71879-2024: классы 5PR/10PR/TPY/TPZ, циклы АПВ (второе КЗ);
- релиз 1.5 — выбор ТТ по ГОСТ Р 71403-2024 (класс точности, I1ном, Kном, Kп.р.ном, Ts ном);
- релиз 1.6 — справочные данные ГОСТ Р 70507.2-2024 (классы, диапазоны параметров).

**Ограничения текущего релиза:**
- только аналитический метод;
- только ТТ классов точности P / PX / TPX с замкнутым магнитопроводом.
""")
